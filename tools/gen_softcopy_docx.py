# -*- coding: utf-8 -*-
"""
生成软著源程序鉴别材料 Word 版（.docx）

与 PDF 版同源同口径（共用 tools/softcopy_common.py）：
  剔除空行/纯注释行 → 超长行按实测宽度折行（不截断）→ 每页 50 行 → 60 页（前 30 + 后 30）

Word 版实现要点：
  - A4 纵向；页边距左右 15mm、上 20mm、下 15mm（与 PDF 版一致）
  - 页眉：左侧「软件全称 + 版本号」，右侧「第 X 页 / 共 60 页」（PAGE / NUMPAGES 域）
  - 正文：每行一个段落，样式固定「行距 = 精确 13.4pt、段前段后 0、取消孤行控制」，
    并在每 50 行后强制分页（第 2~60 页首行设 page_break_before），保证每页恰好 50 行
  - 字体：SimSun（宋体，ASCII 半角 / CJK 全角，度量与折行所用 STSong-Light 一致），7pt

用法：python tools/gen_softcopy_docx.py
输出：dist/taptap/软著源程序_得道飞升模拟器.docx
"""
import os
import sys

from docx import Document
from docx.enum.text import WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from softcopy_common import (FONT_SIZE, LINES_PER_PAGE, LINE_LEAD, MARGIN,  # noqa: E402
                             SOFT_NAME, TAP_DIR, VERSION, fold_lines,
                             load_effective, slice_pages)

OUT_DOCX = os.path.join(TAP_DIR, "软著源程序_得道飞升模拟器.docx")
WORD_FONT = "SimSun"          # 宋体：ASCII 半角、CJK 全角，与折行度量一致
MARGIN_CM = round(MARGIN / 28.3465, 2)     # 15mm → 1.5cm


# ---------------------------------------------------------------- 域代码
def add_page_field(paragraph):
    """在段落中插入「第 {PAGE} 页 / 共 {NUMPAGES} 页」。"""
    def _run(text):
        return paragraph.add_run(text)

    _run("第 ")
    r = paragraph.add_run()
    for tag, attr, val in (("w:fldChar", "w:fldCharType", "begin"),):
        el = OxmlElement(tag)
        el.set(qn(attr), val)
        r._r.append(el)
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    r._r.append(instr)
    sep = OxmlElement("w:fldChar")
    sep.set(qn("w:fldCharType"), "separate")
    r._r.append(sep)
    t = OxmlElement("w:t")
    t.text = "1"                      # 缓存值，Word 打开后自动重算
    r._r.append(t)
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    r._r.append(end)

    _run(" 页 / 共 ")
    r2 = paragraph.add_run()
    b = OxmlElement("w:fldChar")
    b.set(qn("w:fldCharType"), "begin")
    r2._r.append(b)
    i2 = OxmlElement("w:instrText")
    i2.set(qn("xml:space"), "preserve")
    i2.text = " NUMPAGES "
    r2._r.append(i2)
    s2 = OxmlElement("w:fldChar")
    s2.set(qn("w:fldCharType"), "separate")
    r2._r.append(s2)
    t2 = OxmlElement("w:t")
    t2.text = "60"
    r2._r.append(t2)
    e2 = OxmlElement("w:fldChar")
    e2.set(qn("w:fldCharType"), "end")
    r2._r.append(e2)
    _run(" 页")


def _set_font(rpr, name):
    rf = rpr.find(qn("w:rFonts"))
    if rf is None:
        rf = OxmlElement("w:rFonts")
        rpr.append(rf)
    for a in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        rf.set(qn(a), name)


# ---------------------------------------------------------------- 数据
raw, effective = load_effective()
physical = fold_lines(effective)
pages, info = slice_pages(physical)
TOTAL_PAGES = len(pages)

print(f"原始总行数: {len(raw)}")
print(f"有效代码行数: {len(effective)}")
print(f"折行后打印行总数: {len(physical)}")
print(f"页数: {TOTAL_PAGES}"
      + ("" if info["mode"] == "all"
         else f"（前 30 页 = 有效代码 1~{info['cov_front']} 行；"
              f"后 30 页 = {info['total_effective'] - info['cov_back'] + 1}"
              f"~{info['total_effective']} 行）"))

# ---------------------------------------------------------------- 文档
doc = Document()

# 正文样式：等宽、精确行距、无段间距、无孤行控制
style = doc.styles["Normal"]
style.font.name = WORD_FONT
style.font.size = Pt(FONT_SIZE)
_set_font(style.element.get_or_add_rPr(), WORD_FONT)
pf = style.paragraph_format
pf.space_before = Pt(0)
pf.space_after = Pt(0)
pf.line_spacing = Pt(LINE_LEAD)          # 精确行距
pf.widow_control = False

# 关闭网格对齐，避免 Word 按文档网格重新排布导致行数漂移
ppr = style.element.get_or_add_pPr()
snap = OxmlElement("w:snapToGrid")
snap.set(qn("w:val"), "0")
ppr.append(snap)

sec = doc.sections[0]
sec.page_width = Cm(21.0)
sec.page_height = Cm(29.7)
sec.left_margin = Cm(MARGIN_CM)
sec.right_margin = Cm(MARGIN_CM)
sec.top_margin = Cm(2.0)
sec.bottom_margin = Cm(1.5)
sec.header_distance = Cm(1.0)
sec.footer_distance = Cm(1.0)
sec.different_first_page_header_footer = False

# 页眉
hp = sec.header.paragraphs[0]
hp.text = ""
hp.paragraph_format.space_before = Pt(0)
hp.paragraph_format.space_after = Pt(0)
hp.paragraph_format.line_spacing = Pt(10)
hp.paragraph_format.tab_stops.add_tab_stop(Cm(18.0), WD_TAB_ALIGNMENT.RIGHT)
r = hp.add_run(f"{SOFT_NAME} {VERSION}\t")
r.font.size = Pt(8)
r.font.name = WORD_FONT
_set_font(r._r.get_or_add_rPr(), WORD_FONT)
add_page_field(hp)
for run in hp.runs:
    run.font.size = Pt(8)
    run.font.name = WORD_FONT
    _set_font(run._r.get_or_add_rPr(), WORD_FONT)

# 正文：逐行成段，每 50 行分页
n_lines = 0
for page_no, page_lines in enumerate(pages, start=1):
    for i, (txt, _idx) in enumerate(page_lines):
        p = doc.add_paragraph(txt)
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = Pt(LINE_LEAD)
        p.paragraph_format.widow_control = False
        if i == 0 and page_no > 1:
            p.paragraph_format.page_break_before = True
        n_lines += 1

os.makedirs(TAP_DIR, exist_ok=True)
doc.save(OUT_DOCX)
print(f"已生成: {OUT_DOCX}")
print(f"段落数: {len(doc.paragraphs)}（正文 {n_lines} 行 + 分页）/ "
      f"页数: {TOTAL_PAGES}")
print(f"大小: {os.path.getsize(OUT_DOCX)/1024:.1f} KB")
